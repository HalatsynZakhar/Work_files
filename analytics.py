"""
Модуль аналитики, проверки файлов, сопоставления цен конкурентов и формирования сводной таблицы.
"""

import datetime
import json
import os
import re
import shutil
import subprocess
import sys
from decimal import Decimal
from typing import Any, Optional

import openpyxl
from openpyxl.utils import get_column_letter
import pandas as pd

from images_feed import images_feed

if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(PROJECT_ROOT, "data")
INPUT_DIR = os.path.join(DATA_DIR, "input")
OUTPUT_DIR = os.path.join(DATA_DIR, "output")
CACHE_DIR = os.path.join(DATA_DIR, "cache")

os.makedirs(INPUT_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(CACHE_DIR, exist_ok=True)

LOCAL_CACHE_WATCHES = os.path.join(CACHE_DIR, "url-watches.json")
EXTERNAL_BACKUP_DIR = r"C:\Users\ABM\Desktop\Робота\Звіти керівництву\27 Подтягивая цен за мониторингом на Хорошоп"

CONTAINER_NAMES = [
    "changedetectionio-changedetection-1",
    "changedetection",
    "changedetectionio_changedetection_1"
]
REMOTE_JSON_PATH = "/datastore/url-watches.json"
BRAND_NAMES = {'road rippers', 'syma', 'nikko', 'revolt', 'machine maker'}
PERSONAL_TAGS = {'личное', 'личные', 'personal', 'private'}

DOCKER_LIVE_EXTRACT_SCRIPT = """
import os, json, glob

datastore = "/datastore"
watches = {}
for p in glob.glob(os.path.join(datastore, "*", "watch.json")):
    uuid = os.path.basename(os.path.dirname(p))
    try:
        with open(p, "r", encoding="utf-8") as f:
            watches[uuid] = json.load(f)
    except Exception:
        pass

tags = {}
for p in glob.glob(os.path.join(datastore, "*", "tag.json")):
    uuid = os.path.basename(os.path.dirname(p))
    try:
        with open(p, "r", encoding="utf-8") as f:
            tags[uuid] = json.load(f)
    except Exception:
        pass

settings = {}
cd_json = os.path.join(datastore, "changedetection.json")
if os.path.exists(cd_json):
    try:
        with open(cd_json, "r", encoding="utf-8") as f:
            settings = json.load(f).get("settings", {})
    except Exception:
        pass

if "application" not in settings:
    settings["application"] = {}
settings["application"]["tags"] = tags

print(json.dumps({"watching": watches, "settings": settings}, ensure_ascii=False))
"""

INPUT_FILES_SPEC = [
    {
        "key": "catalog",
        "title": "Каталог артикулов (базовый файл)",
        "preferred": "articles_catalog.xlsm",
        "alias": "articles.xlsm",
        "required": True,
        "max_age_days": 14,
        "description": "Содержит соответствие Артикул <-> ХОРОШОП АРТИКУЛ"
    },
    {
        "key": "rrc",
        "title": "Таблица РРЦ (Рекомендованные розничные цены)",
        "preferred": "rrc_prices.xlsx",
        "alias": "РРЦ.xlsx",
        "required": True,
        "max_age_days": 7,
        "description": "Цены РРЦ для расчета скидок опта и анализа демпинга"
    },
    {
        "key": "locked",
        "title": "Таблица зафиксированных цен (Замки UI)",
        "preferred": "locked_articles.xlsx",
        "alias": "locked_articles.json",
        "required": False,
        "max_age_days": 365,
        "description": "Управляется интерактивным замочком в веб-интерфейсе"
    }
]


def resolve_file_path(preferred: str, alias: str) -> Optional[str]:
    """Ищет файл сначала по предпочтительному имени, затем по алиасу."""
    candidates = [
        os.path.join(INPUT_DIR, preferred),
        os.path.join(INPUT_DIR, alias),
        os.path.join(DATA_DIR, preferred),
        os.path.join(DATA_DIR, alias),
        os.path.join(PROJECT_ROOT, preferred),
        os.path.join(PROJECT_ROOT, alias),
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return None


def check_startup_files() -> list[dict[str, Any]]:
    """
    Проверяет наличие и дату изменения ключевых файлов в data/input.
    Формирует список предупреждений, если файл устарел или отсутствует.
    """
    results = []
    now = datetime.datetime.now()

    for spec in INPUT_FILES_SPEC:
        path = resolve_file_path(spec["preferred"], spec["alias"])
        if not path:
            results.append({
                "key": spec["key"],
                "title": spec["title"],
                "filename": spec["preferred"],
                "exists": False,
                "path": os.path.join(INPUT_DIR, spec["preferred"]),
                "modified": "",
                "age_days": None,
                "is_warning": True,
                "warning_level": "danger" if spec["required"] else "warning",
                "message": f"Файл '{spec['preferred']}' не найден в data/input/! ({spec['description']})"
            })
            continue

        mtime = datetime.datetime.fromtimestamp(os.path.getmtime(path))
        age_days = (now - mtime).total_seconds() / 86400.0
        max_age = spec["max_age_days"]

        is_stale = age_days > max_age
        msg = f"Актуален (обновлен {mtime.strftime('%d.%m.%Y %H:%M')})"
        warning_level = "success"

        if is_stale:
            msg = (
                f"Файл не обновлялся {int(age_days)} дн. (предел: {max_age} дн.). "
                f"Рекомендуется проверить актуальность перед выгрузкой цен!"
            )
            warning_level = "warning"

        results.append({
            "key": spec["key"],
            "title": spec["title"],
            "filename": os.path.basename(path),
            "exists": True,
            "path": path,
            "modified": mtime.strftime("%d.%m.%Y %H:%M"),
            "age_days": round(age_days, 1),
            "is_warning": is_stale,
            "warning_level": warning_level,
            "message": msg
        })

    return results


LOCKED_JSON_PATH = os.path.join(DATA_DIR, "locked_articles.json")
LOCKED_XLSX_PATH = os.path.join(DATA_DIR, "locked_articles.xlsx")


def load_locked_articles() -> set[str]:
    """
    Загружает список заблокированных артикулов из собственной таблицы системы:
    1. Проверяет data/locked_articles.json.
    2. Если его нет, проверяет data/locked_articles.xlsx.
    3. Если и его нет, импортирует начальные позиции из fixed_locked_articles.xlsx (если есть).
    """
    if os.path.exists(LOCKED_JSON_PATH):
        try:
            with open(LOCKED_JSON_PATH, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, list):
                return {str(x).strip().upper() for x in data if str(x).strip()}
            elif isinstance(data, dict):
                return {str(k).strip().upper() for k, v in data.items() if v}
        except Exception as e:
            print(f"[WARN] Ошибка чтения {LOCKED_JSON_PATH}: {e}")

    if os.path.exists(LOCKED_XLSX_PATH):
        try:
            wb = openpyxl.load_workbook(LOCKED_XLSX_PATH, read_only=True, data_only=True)
            ws = wb.active
            locked = set()
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None:
                    val = str(row[0]).strip().upper()
                    if val and val != "NAN":
                        locked.add(val)
            wb.close()
            # Сохраняем в JSON для быстрого доступа
            save_locked_articles(locked)
            return locked
        except Exception as e:
            print(f"[WARN] Ошибка чтения {LOCKED_XLSX_PATH}: {e}")

    # Первичная миграция из старого пользовательского файла fixed_locked_articles.xlsx если он есть
    legacy_path = resolve_file_path("fixed_locked_articles.xlsx", "ФИКСИРОВАТЬ ЦЕНУ.xlsx")
    if legacy_path and os.path.exists(legacy_path):
        try:
            wb = openpyxl.load_workbook(legacy_path, read_only=True, data_only=True)
            ws = wb.active
            locked = set()
            for row in ws.iter_rows(values_only=True):
                if not row or row[0] is None:
                    continue
                val = str(row[0]).strip().upper()
                if val and val not in ("АРТИКУЛ", "ARTICLE", "NAN"):
                    locked.add(val)
            wb.close()
            save_locked_articles(locked)
            print(f"[OK] Мигрировано {len(locked)} заблокированных артикулов из {os.path.basename(legacy_path)} в таблицу замков.")
            return locked
        except Exception as e:
            print(f"[WARN] Ошибка первичной миграции locked_articles: {e}")

    return set()


def save_locked_articles(locked_set: set[str]) -> None:
    """
    Сохраняет список заблокированных артикулов в собственную JSON базу и Excel-таблицу.
    """
    clean_list = sorted(list({str(x).strip().upper() for x in locked_set if str(x).strip()}))

    # 1. Сохранение в JSON
    try:
        tmp_json = LOCKED_JSON_PATH + ".tmp"
        with open(tmp_json, "w", encoding="utf-8") as f:
            json.dump(clean_list, f, ensure_ascii=False, indent=2)
        if os.path.exists(LOCKED_JSON_PATH):
            os.remove(LOCKED_JSON_PATH)
        os.rename(tmp_json, LOCKED_JSON_PATH)
    except Exception as e:
        print(f"[ERR] Ошибка сохранения {LOCKED_JSON_PATH}: {e}")

    # 2. Сохранение в Excel-таблицу locked_articles.xlsx для наглядности пользователю
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Зафиксированные цены"
        ws.append(["Артикул", "Статус", "Зафиксировано в системе"])
        now_str = datetime.datetime.now().strftime("%d.%m.%Y %H:%M")
        for art in clean_list:
            ws.append([art, "ЗАФИКСИРОВАНО", now_str])

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 25

        # Стилизация заголовка
        from openpyxl.styles import Font, PatternFill, Alignment
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="C0392B")
            cell.alignment = Alignment(horizontal="center")

        tmp_xlsx = LOCKED_XLSX_PATH + ".tmp"
        wb.save(tmp_xlsx)
        wb.close()
        if os.path.exists(LOCKED_XLSX_PATH):
            os.remove(LOCKED_XLSX_PATH)
        os.rename(tmp_xlsx, LOCKED_XLSX_PATH)
    except Exception as e:
        print(f"[WARN] Ошибка сохранения {LOCKED_XLSX_PATH}: {e}")


def toggle_article_lock(article: str, locked_state: Optional[bool] = None) -> bool:
    """
    Переключает или устанавливает статус замка (блокировки) для артикула.
    Возвращает новый статус: True если заблокирован, False если разблокирован.
    """
    art_clean = str(article).strip().upper()
    if not art_clean:
        return False

    current_locked = load_locked_articles()

    if locked_state is not None:
        new_state = bool(locked_state)
    else:
        new_state = art_clean not in current_locked

    if new_state:
        current_locked.add(art_clean)
    else:
        current_locked.discard(art_clean)

    save_locked_articles(current_locked)
    return new_state


def load_rrc_prices() -> dict[str, float]:
    """
    Загружает справочник РРЦ из rrc_prices.xlsx.
    Использует data_only=True для чтения кэшированных значений формул XLOOKUP.
    """
    path = resolve_file_path("rrc_prices.xlsx", "РРЦ.xlsx")
    if not path or not os.path.exists(path):
        return {}

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        rrc_map: dict[str, float] = {}

        # Чтение заголовка
        headers = [str(cell.value or "").strip().lower() for cell in ws[1]]
        art_idx = 0
        rrc_idx = 1
        for idx, h in enumerate(headers):
            if "артикул" in h or "article" in h:
                art_idx = idx
            elif "ррц" in h or "ціна" in h or "цена" in h or "price" in h:
                rrc_idx = idx

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) <= max(art_idx, rrc_idx):
                continue
            art_val = str(row[art_idx] or "").strip().upper()
            price_val = row[rrc_idx]
            if not art_val or art_val == "NAN":
                continue

            try:
                if price_val is not None:
                    p_float = float(str(price_val).replace(" ", "").replace(",", "."))
                    if p_float > 0:
                        rrc_map[art_val] = p_float
            except Exception:
                continue

        wb.close()
        return rrc_map
    except Exception as e:
        print(f"[WARN] Ошибка загрузки rrc_prices: {e}")
        return {}


def load_catalog_dataframe() -> pd.DataFrame:
    """
    Загружает базу соответствий артикулов и GUID Хорошоп (Каталог товаров).
    Считывает листы 'Input data' (номенклатура 1C/Хорошоп, характеристики)
    и 'Python data' (прямые сопоставления Артикул <-> ХОРОШОП АРТИКУЛ / GUID).
    """
    path = resolve_file_path("articles_catalog.xlsm", "articles.xlsm")
    if not path or not os.path.exists(path):
        return pd.DataFrame(columns=["Артикул", "ХОРОШОП АРТИКУЛ", "Название"])

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        mappings: dict[str, dict[str, str]] = {}

        # 1. Лист 'Input data' (исходные данные 1C/Хорошоп: номенклатура, артикулы модификаций и GUID)
        if 'Input data' in wb.sheetnames:
            for r in wb['Input data'].iter_rows(values_only=True):
                if not r or not r[0]:
                    continue
                art = str(r[0]).strip()
                if art.lower() in ['артикул', 'none', 'nan', '']:
                    continue
                name = str(r[1]).strip() if len(r) > 1 and r[1] else ''
                art_char = str(r[2]).strip() if len(r) > 2 and r[2] else ''
                guid = str(r[4]).strip() if len(r) > 4 and r[4] else ''
                if guid.lower() in ['характеристика', 'none', 'nan', '-']:
                    guid = ''

                # Базовый артикул
                if art not in mappings:
                    mappings[art] = {'guid': guid, 'name': name}
                else:
                    if guid and not mappings[art]['guid']:
                        mappings[art]['guid'] = guid
                    if name and not mappings[art]['name']:
                        mappings[art]['name'] = name

                # Артикул модификации / цвета (например X15A White, TG1009 BLUE)
                if art_char and art_char.lower() not in ['none', 'nan', '']:
                    if art_char not in mappings:
                        mappings[art_char] = {'guid': guid, 'name': name}
                    elif guid and not mappings[art_char]['guid']:
                        mappings[art_char]['guid'] = guid

        # 2. Лист 'Python data' (сопоставления Артикул <-> GUID)
        if 'Python data' in wb.sheetnames:
            for r in wb['Python data'].iter_rows(values_only=True):
                if not r or not r[0]:
                    continue
                art = str(r[0]).strip()
                guid = str(r[1]).strip() if len(r) > 1 and r[1] else ''
                if art.lower() in ['артикул', 'none', 'nan', '']:
                    continue
                if guid.lower() in ['характеристика', 'хорошоп артикул', 'none', 'nan', '-']:
                    guid = ''

                if art:
                    if art not in mappings:
                        mappings[art] = {'guid': guid, 'name': ''}
                    elif guid:
                        mappings[art]['guid'] = guid

        if not mappings:
            df = pd.read_excel(path, sheet_name=0, engine="openpyxl")
            col_map = {}
            for c in df.columns:
                c_str = str(c).strip()
                c_lower = c_str.lower()
                if "хорошоп" in c_lower or "uuid" in c_lower or "guid" in c_lower:
                    col_map[c] = "ХОРОШОП АРТИКУЛ"
                elif "артикул" in c_lower or "article" in c_lower:
                    col_map[c] = "Артикул"
                elif "назв" in c_lower or "номенклатура" in c_lower:
                    col_map[c] = "Название"

            df = df.rename(columns=col_map)
            if "Артикул" not in df.columns and len(df.columns) > 0:
                df = df.rename(columns={df.columns[0]: "Артикул"})
            if "ХОРОШОП АРТИКУЛ" not in df.columns and len(df.columns) > 1:
                df = df.rename(columns={df.columns[1]: "ХОРОШОП АРТИКУЛ"})
            if "Название" not in df.columns:
                df["Название"] = ""

            df["Артикул"] = df["Артикул"].astype(str).str.strip()
            df["ХОРОШОП АРТИКУЛ"] = df["ХОРОШОП АРТИКУЛ"].astype(str).str.strip().replace({"nan": "", "None": ""})
            df = df[~df["Артикул"].str.lower().isin(["артикул", "none", "nan", ""])]
            return df

        rows = []
        for art, info in mappings.items():
            rows.append({
                "Артикул": art,
                "ХОРОШОП АРТИКУЛ": info["guid"],
                "Название": info["name"]
            })
        return pd.DataFrame(rows)
    except Exception as e:
        print(f"[WARN] Ошибка загрузки каталога артикулов: {e}")
        return pd.DataFrame(columns=["Артикул", "ХОРОШОП АРТИКУЛ", "Название"])


def find_running_container() -> Optional[str]:
    """Находит имя активного контейнера changedetection."""
    try:
        cmd = ["docker", "ps", "--filter", "ancestor=dgtlmoon/changedetection.io", "--format", "{{.Names}}"]
        res = subprocess.run(cmd, capture_output=True, text=True, check=True)
        containers = [c.strip() for c in res.stdout.splitlines() if c.strip()]
        if containers:
            return containers[0]

        cmd = ["docker", "ps", "--filter", "name=changedetection", "--format", "{{.Names}}"]
        res = subprocess.run(cmd, capture_output=True, text=True, check=True)
        containers = [c.strip() for c in res.stdout.splitlines() if c.strip()]
        if containers:
            filtered = [c for c in containers if "selenium" not in c.lower()]
            if filtered:
                return filtered[0]

        for name in CONTAINER_NAMES:
            check_cmd = ["docker", "inspect", "-f", "{{.State.Running}}", name]
            r = subprocess.run(check_cmd, capture_output=True, text=True)
            if r.returncode == 0 and r.stdout.strip() == "true":
                return name
    except Exception:
        pass
    return None


def get_docker_watches_data(force_refresh: bool = False) -> tuple[dict[str, Any], bool]:
    """
    Получает данные из Docker контейнера changedetection.io.
    Извлекает активные watches и теги напрямую из директорий /datastore.
    Кэширует в data/cache/url-watches.json.
    Возвращает: (данные_json, is_live_docker: bool)
    """
    container = find_running_container()
    if container:
        # 1. Приоритет: получение актуальных watches напрямую из директорий /datastore/*/watch.json
        try:
            cmd = ["docker", "exec", container, "python3", "-c", DOCKER_LIVE_EXTRACT_SCRIPT]
            res = subprocess.run(cmd, capture_output=True, check=True)
            raw = res.stdout.decode("utf-8", errors="ignore")
            data = json.loads(raw)
            if data and data.get("watching"):
                with open(LOCAL_CACHE_WATCHES, "w", encoding="utf-8") as f:
                    f.write(raw)
                return data, True
        except Exception as e:
            print(f"[WARN] Ошибка извлечения активных watches из директорий Docker: {e}")

        # 2. Fallback: чтение url-watches.json внутри контейнера
        try:
            cmd = ["docker", "exec", container, "cat", REMOTE_JSON_PATH]
            res = subprocess.run(cmd, capture_output=True, check=True)
            raw = res.stdout.decode("utf-8", errors="ignore")
            data = json.loads(raw)
            with open(LOCAL_CACHE_WATCHES, "w", encoding="utf-8") as f:
                f.write(raw)
            return data, True
        except Exception as e:
            print(f"[WARN] Не удалось прочитать url-watches.json из Docker: {e}")

    # 3. Fallback на локальный кэш
    if os.path.exists(LOCAL_CACHE_WATCHES):
        try:
            with open(LOCAL_CACHE_WATCHES, "r", encoding="utf-8") as f:
                data = json.load(f)
            return data, False
        except Exception:
            pass

    return {}, False


def extract_article_info(url: str, title: str = "") -> tuple[str, str, str]:
    """Извлекает артикул из URL или названия товара."""
    article_match = re.search(r'/([^/]+)\.html', url)
    raw_article = article_match.group(1) if article_match else ""
    art_from_url = raw_article.split('-')[-1].strip().upper() if raw_article else ""

    art_from_title = ""
    if title:
        parts = title.strip().split()
        if parts:
            art_from_title = parts[0].upper()

    return raw_article, art_from_url, art_from_title


def parse_all_watches(data: dict[str, Any]) -> list[dict[str, Any]]:
    """Парсит все watches из Changedetection в плоский структурированный список."""
    watching = data.get("watching", {})
    tags_dict = {
        k: v.get("title", k)
        for k, v in data.get("settings", {}).get("application", {}).get("tags", {}).items()
    }

    records = []
    for uuid, w in watching.items():
        restock = w.get("restock") or {}
        if w.get("processor") != "restock_diff" and not restock:
            continue

        url = w.get("url", "")
        title = (w.get("title") or "").strip()
        price = restock.get("price")
        in_stock = bool(restock.get("in_stock", False))
        currency = restock.get("currency") or "UAH"

        raw_art, art_url, art_title = extract_article_info(url, title)

        w_tag_ids = w.get("tags", [])
        tag_names = [tags_dict.get(t, t) for t in w_tag_ids]
        is_personal = any(t.lower() in PERSONAL_TAGS for t in tag_names)
        stores = [t for t in tag_names if t.lower() not in BRAND_NAMES and t.lower() not in PERSONAL_TAGS]
        brands = [t for t in tag_names if t.lower() in BRAND_NAMES]

        last_checked = ""
        ts = w.get("last_checked")
        if ts:
            try:
                last_checked = datetime.datetime.fromtimestamp(int(ts)).strftime("%Y-%m-%d %H:%M")
            except Exception:
                last_checked = str(ts)

        records.append({
            "uuid": uuid,
            "url": url,
            "title": title,
            "title_upper": title.upper(),
            "raw_article": raw_art.upper(),
            "art_from_url": art_url,
            "art_from_title": art_title,
            "price": float(price) if price is not None else None,
            "currency": currency,
            "in_stock": in_stock,
            "store": ", ".join(stores) if stores else ("Личное" if is_personal else "Другие магазины"),
            "brand": ", ".join(brands) if brands else "",
            "last_checked": last_checked,
            "is_personal": is_personal
        })

    return records


def build_consolidated_pivot(
    df_catalog: pd.DataFrame,
    watches_list: list[dict[str, Any]],
    rrc_map: dict[str, float],
    locked_articles: set[str],
    horoshop_catalog: Optional[dict[str, dict[str, Any]]] = None
) -> list[dict[str, Any]]:
    """
    Строит сводную аналитическую таблицу:
    1 строка = 1 артикул, с агрегацией всех предложений рынка, РРЦ, ценой на сайте и расчетом минимальной цены.
    """
    images_feed.ensure_cache()
    if not images_feed._loaded:
        images_feed.load()

    rows = []
    for _, cat_row in df_catalog.iterrows():
        my_art = str(cat_row.get("Артикул", "")).strip().upper()
        if not my_art or my_art in ["NAN", "NONE", "АРТИКУЛ"]:
            continue

        horoshop_art = str(cat_row.get("ХОРОШОП АРТИКУЛ", "")).strip()
        if horoshop_art.upper() in ["ХАРАКТЕРИСТИКА", "ХОРОШОП АРТИКУЛ", "NONE", "NAN", "-"]:
            horoshop_art = ""
        is_locked = (my_art in locked_articles) or (bool(horoshop_art) and horoshop_art.upper() in locked_articles)

        # Поиск всех предложений конкурентов для данного артикула
        matched_offers = []
        # Шаблон точного поиска артикула в названии (защита от склеек, например CNPS9X не совпадет с S9)
        art_pattern = re.compile(r'(?<![A-Za-z0-9])' + re.escape(my_art) + r'(?![A-Za-z0-9])', re.IGNORECASE)
        # Шаблон поиска артикула в URL (только как отдельный сегмент между разделителями / _ - . )
        url_art_pattern = re.compile(r'(?:^|[_\-/.])' + re.escape(my_art) + r'(?:[_\-/.]|$)', re.IGNORECASE)

        for w in watches_list:
            # Игнорируем личные отслеживания
            if w.get("is_personal"):
                continue

            # Очищаем URL от параметров запроса (?gad_source=..., ?gclid=... и т.д.), чтобы случайные токены не триггерили совпадение
            clean_url = w["url"].split('?')[0].split('#')[0]

            is_match = False
            if w["art_from_url"] == my_art:
                is_match = True
            elif art_pattern.search(w["title_upper"]):
                is_match = True
            elif len(my_art) > 3 and url_art_pattern.search(clean_url):
                is_match = True
            elif len(my_art) <= 3 and url_art_pattern.search(clean_url):
                # Для сверхкоротких артикулов (<= 3 символов, напр. S9, Z4) требуется подтверждение в названии или точном raw_article
                if w["raw_article"] == my_art or art_pattern.search(w["title_upper"]):
                    is_match = True

            if is_match:
                matched_offers.append(w)

        # Выделение цен в наличии
        in_stock_prices = [
            o["price"] for o in matched_offers
            if o["in_stock"] and o["price"] is not None and o["price"] > 0
        ]
        all_prices = [
            o["price"] for o in matched_offers
            if o["price"] is not None and o["price"] > 0
        ]

        min_market = min(in_stock_prices) if in_stock_prices else (min(all_prices) if all_prices else None)
        max_market = max(in_stock_prices) if in_stock_prices else (max(all_prices) if all_prices else None)
        avg_market = round(sum(in_stock_prices) / len(in_stock_prices), 2) if in_stock_prices else (
            round(sum(all_prices) / len(all_prices), 2) if all_prices else None
        )

        min_store = ""
        if min_market is not None:
            for o in matched_offers:
                if o["price"] == min_market and (o["in_stock"] or not in_stock_prices):
                    min_store = o["store"]
                    break

        # РРЦ
        rrc_price = rrc_map.get(my_art) or (rrc_map.get(horoshop_art.upper()) if horoshop_art else None)

        # Наша цена на Хорошоп (из выгрузки API)
        our_price = None
        our_price_old = None
        p_info = None
        if horoshop_catalog:
            # 1. Приоритет: точный поиск по ХОРОШОП АРТИКУЛ (внутренний UUID Хорошоп вида ce06d003-8aa5-11ed-870e-1402ec4177b7)
            if horoshop_art:
                p_info = horoshop_catalog.get(horoshop_art.upper())

            # 2. Если не найден по UUID, поиск по display_article (my_art)
            if not p_info:
                p_info = horoshop_catalog.get(my_art)

            # 3. Поиск по нормализованному коду (без пробелов, дефисов и спецсимволов)
            if not p_info:
                my_art_clean = re.sub(r'[^A-Z0-9]', '', my_art)
                for h_k, h_v in horoshop_catalog.items():
                    if re.sub(r'[^A-Z0-9]', '', h_k) == my_art_clean:
                        p_info = h_v
                        break

            # Если в каталоге найден товар, но horoshop_art не был указан в файле, берем internal UUID из каталога
            if p_info and not horoshop_art:
                horoshop_art = str(p_info.get("article", "")).strip()

            if p_info:
                raw_p = p_info.get("price")
                raw_old = p_info.get("price_old")
                if raw_p is not None:
                    try:
                        our_price = float(raw_p)
                    except (ValueError, TypeError):
                        pass
                if raw_old is not None:
                    try:
                        our_price_old = float(raw_old)
                    except (ValueError, TypeError):
                        pass

                if not rrc_price and our_price_old and our_price_old > 0:
                    rrc_price = our_price_old

        horoshop_brand = p_info.get("brand_title", "") if p_info else ""

        # Изображение, бренд и название
        img_url = images_feed.get_image(my_art)
        feed_brand = images_feed.get_brand(my_art)
        title = images_feed.get_name(my_art)
        catalog_title = str(cat_row.get("Название", "")).strip()

        # Приоритет бренда: 1. Хорошоп (из API каталога), 2. XML фид, 3. Теги трекера
        brand = horoshop_brand or feed_brand or ""
        if not brand:
            for o in matched_offers:
                if o.get("brand"):
                    brand = o["brand"]
                    break

        # Если в фиде нет названия, берем официальное из каталога или от партнера
        if not title:
            title = catalog_title or (matched_offers[0]["title"] if matched_offers else "")

        # Партнерские цены по магазинам
        partner_breakdown: dict[str, dict[str, Any]] = {}
        for o in matched_offers:
            st = o["store"]
            # Сохраняем лучшее предложение магазина
            if st not in partner_breakdown or (o["in_stock"] and not partner_breakdown[st]["in_stock"]):
                partner_breakdown[st] = {
                    "price": o["price"],
                    "in_stock": o["in_stock"],
                    "url": o["url"],
                    "last_checked": o["last_checked"]
                }

        # 1. Анализ отклонения от РРЦ (Рынок vs РРЦ)
        diff_to_rrc = round(min_market - rrc_price, 2) if (min_market is not None and rrc_price is not None) else None
        diff_to_rrc_pct = round(((min_market - rrc_price) / rrc_price) * 100, 1) if (min_market is not None and rrc_price and rrc_price > 0) else None

        # 2. КРИТИЧЕСКОЕ ПРАВИЛО ДИСТРИБЬЮТОРА:
        # Отклонение от нашего сайта (Наша цена vs Мин. рынок)
        diff_our_to_market = round(our_price - min_market, 2) if (our_price is not None and min_market is not None) else None
        diff_our_to_market_pct = round(((our_price - min_market) / min_market) * 100, 1) if (our_price is not None and min_market and min_market > 0) else None

        is_our_price_lowest = False
        if our_price is not None and min_market is not None:
            if our_price < (min_market - 0.5):
                is_our_price_lowest = True

        # Рекомендуемая цена
        recommended_price = min_market if min_market is not None else (rrc_price if rrc_price else our_price)

        rows.append({
            "article": my_art,
            "horoshop_article": horoshop_art,
            "title": title,
            "brand": brand,
            "horoshop_brand": horoshop_brand,
            "image_url": img_url,
            "is_locked": is_locked,
            "rrc_price": rrc_price,
            "our_price": our_price,
            "our_price_old": our_price_old,
            "min_market_price": min_market,
            "min_store": min_store,
            "max_market_price": max_market,
            "avg_market_price": avg_market,
            "diff_to_rrc": diff_to_rrc,
            "diff_to_rrc_pct": diff_to_rrc_pct,
            "diff_our_to_min": diff_our_to_market,
            "diff_our_to_market": diff_our_to_market,
            "diff_our_to_market_pct": diff_our_to_market_pct,
            "is_our_price_lowest": is_our_price_lowest,
            "offers_count": len(matched_offers),
            "in_stock_count": len(in_stock_prices),
            "offers": matched_offers,
            "partner_breakdown": partner_breakdown,
            "recommended_price": recommended_price
        })

    return rows


def auto_fit_excel_columns(excel_path: str) -> None:
    """Автоматическая настройка ширины колонок в Excel файле."""
    try:
        wb = openpyxl.load_workbook(excel_path)
        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = 0
                for cell in col:
                    val = str(cell.value or "")
                    if len(val) > max_len:
                        max_len = len(val)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)
        wb.save(excel_path)
    except Exception as e:
        print(f"[NOTE] Авто-ширина колонок: {e}")


def save_excel_safely(writer_callback, target_path: str, label: str) -> tuple[Optional[str], bool]:
    """Безопасное сохранение Excel с созданием копии при открытом файле."""
    try:
        with pd.ExcelWriter(target_path, engine="openpyxl") as writer:
            writer_callback(writer)
        auto_fit_excel_columns(target_path)
        return target_path, False
    except PermissionError:
        now_str = datetime.datetime.now().strftime("%H%M%S")
        dir_name = os.path.dirname(target_path)
        base_name, ext = os.path.splitext(os.path.basename(target_path))
        fallback_path = os.path.join(dir_name, f"{base_name}_{now_str}{ext}")
        try:
            with pd.ExcelWriter(fallback_path, engine="openpyxl") as writer:
                writer_callback(writer)
            auto_fit_excel_columns(fallback_path)
            print(f"[WARN] Файл '{os.path.basename(target_path)}' открыт в Excel! Сохранен как: {os.path.basename(fallback_path)}")
            return fallback_path, True
        except Exception as e:
            print(f"[ERR] Не удалось сохранить файл: {e}")
            return None, False
    except Exception as e:
        print(f"[ERR] Ошибка сохранения {label}: {e}")
        return None, False


def generate_full_excel_report(
    consolidated_rows: list[dict[str, Any]],
    raw_watches: list[dict[str, Any]],
    chosen_store: Optional[str] = None,
    strategy: Optional[str] = None,
    custom_prices: Optional[dict[str, float]] = None,
    columns_to_export: Optional[list[str]] = None
) -> Optional[str]:
    """
    Формирует комплексный Excel-отчет:
    1. Лист 'Готово к импорту' (настраиваемый состав колонок и цен по выбранной стратегии или ручным полям)
    2. Лист 'Сводная аналитика' (1 артикул = 1 строка с min/max/avg, РРЦ, разницей и ценами магазинов)
    3. Лист 'Детальное сопоставление' (все сопоставленные ссылки и предложения)
    4. Лист 'Вся база трекера Docker'
    """
    output_path = os.path.join(OUTPUT_DIR, "ИТОГОВЫЙ_ОТЧЕТ.xlsx")

    # 1. Лист 'Готово к импорту' (только не заблокированные, с итоговой ценой > 0)
    sheet1_rows = []
    for r in consolidated_rows:
        if r["is_locked"]:
            continue

        art = r["article"]
        hor_art = r.get("horoshop_article", "")

        price_to_use = None

        # 1.1. Наивысший приоритет: значения из переданных полей таблицы (custom_prices)
        if custom_prices:
            if art in custom_prices and custom_prices[art] is not None and float(custom_prices[art]) > 0:
                price_to_use = float(custom_prices[art])
            elif hor_art and hor_art in custom_prices and custom_prices[hor_art] is not None and float(custom_prices[hor_art]) > 0:
                price_to_use = float(custom_prices[hor_art])

        # 1.2. Если в custom_prices цена не задана, берем согласно выбранной стратегии
        if price_to_use is None:
            target_store = chosen_store
            if not target_store and strategy and strategy.startswith("partner_"):
                target_store = strategy.replace("partner_", "").strip()

            if target_store:
                p_entry = r.get("partner_breakdown", {}).get(target_store)
                if p_entry and p_entry.get("price") and p_entry["price"] > 0:
                    price_to_use = float(p_entry["price"])
                else:
                    price_to_use = r.get("min_market_price") or r.get("rrc_price") or r.get("our_price")
            elif strategy == "rrc":
                price_to_use = r.get("rrc_price") or r.get("our_price") or r.get("min_market_price")
            elif strategy == "current":
                price_to_use = r.get("our_price") or r.get("rrc_price") or r.get("min_market_price")
            else:
                # По умолчанию (или min_in_stock)
                price_to_use = r.get("recommended_price") or r.get("min_market_price") or r.get("rrc_price") or r.get("our_price")

        if price_to_use and price_to_use > 0:
            full_record = {
                "Артикул": art,
                "ХОРОШОП АРТИКУЛ": hor_art,
                "Название": r.get("title", ""),
                "Бренд": r.get("brand", ""),
                "Цена": float(price_to_use),
                "Старая цена": r.get("rrc_price") or r.get("our_price_old") or ""
            }

            if columns_to_export and isinstance(columns_to_export, list) and len(columns_to_export) > 0:
                # Отбираем только запрошенные пользователем поля, сохраняя порядок
                row_filtered = {col: full_record.get(col, "") for col in columns_to_export if col in full_record}
                sheet1_rows.append(row_filtered if row_filtered else full_record)
            else:
                # По умолчанию: стандартные 3 колонки для Хорошопа
                sheet1_rows.append({
                    "Артикул": art,
                    "ХОРОШОП АРТИКУЛ": hor_art,
                    "Цена": float(price_to_use)
                })

    df_sheet1 = pd.DataFrame(sheet1_rows)

    # 2. Лист 'Сводная аналитика'
    sheet2_rows = []
    # Соберем уникальные магазины
    all_stores = sorted(list({o["store"] for r in consolidated_rows for o in r["offers"] if o.get("store")}))

    for r in consolidated_rows:
        status_str = "🔒 ЗАФИКСИРОВАНО" if r["is_locked"] else ("В наличии" if r["in_stock_count"] > 0 else "Нет на рынке")

        distributor_note = "Норма (в рынке)"
        if r.get("is_our_price_lowest"):
            diff_abs = abs(r["diff_our_to_min"]) if r["diff_our_to_min"] is not None else 0
            distributor_note = f"⚠️ ДЕМПИНГ ДИЛЕРОВ! (Наша цена ниже на {diff_abs} грн, чем у {r['min_store']})"
        elif r.get("our_price") and r.get("min_market_price") and abs(r["our_price"] - r["min_market_price"]) < 0.5:
            distributor_note = "✓ Цена установлена по мин. рынку"
        elif r.get("our_price") and r.get("min_market_price") and r["our_price"] > r["min_market_price"]:
            distributor_note = f"Выше рынка (+{r['diff_our_to_min']} грн)"

        row_dict = {
            "Артикул": r["article"],
            "ХОРОШОП АРТИКУЛ": r["horoshop_article"],
            "Бренд": r["brand"],
            "Название товара": r["title"],
            "Статус": status_str,
            "Статус дистрибьютора": distributor_note,
            "РРЦ (грн)": r["rrc_price"],
            "Наша цена на сайте": r["our_price"],
            "Мин. цена на рынке": r["min_market_price"],
            "Магазин мин. цены": r["min_store"],
            "Макс. цена на рынке": r["max_market_price"],
            "Средняя цена на рынке": r["avg_market_price"],
            "Отклонение от РРЦ (грн)": r.get("diff_to_rrc"),
            "Отклонение от РРЦ (%)": r.get("diff_to_rrc_pct"),
            "Отклонение от Нас (грн)": r.get("diff_our_to_market"),
            "Отклонение от Нас (%)": r.get("diff_our_to_market_pct"),
            "Рекомендуемая цена": r["recommended_price"]
        }
        # Цены конкретных магазинов
        for st in all_stores:
            p_info = r["partner_breakdown"].get(st)
            if p_info:
                in_stk_mark = " (в наличии)" if p_info["in_stock"] else " (нет)"
                row_dict[st] = f"{p_info['price']} грн{in_stk_mark}"
            else:
                row_dict[st] = ""

        sheet2_rows.append(row_dict)

    df_sheet2 = pd.DataFrame(sheet2_rows)

    # 3. Лист 'Детальное сопоставление'
    sheet3_rows = []
    for r in consolidated_rows:
        for o in r["offers"]:
            sheet3_rows.append({
                "Артикул": r["article"],
                "ХОРОШОП АРТИКУЛ": r["horoshop_article"],
                "Магазин": o["store"],
                "Цена партнера": o["price"],
                "В наличии": "Да" if o["in_stock"] else "Нет",
                "Название у партнера": o["title"],
                "Ссылка на товар": o["url"],
                "Дата проверки": o["last_checked"]
            })
    df_sheet3 = pd.DataFrame(sheet3_rows)

    # 4. Лист 'Вся база трекера Docker'
    df_sheet4 = pd.DataFrame(raw_watches)
    if not df_sheet4.empty:
        col_order = ["title", "store", "brand", "price", "in_stock", "url", "last_checked"]
        existing_cols = [c for c in col_order if c in df_sheet4.columns]
        df_sheet4 = df_sheet4[existing_cols].rename(columns={
            "title": "Название товара",
            "store": "Магазин",
            "brand": "Бренд",
            "price": "Цена",
            "in_stock": "В наличии",
            "url": "Ссылка",
            "last_checked": "Дата проверки"
        })

    def write_workbook(writer):
        df_sheet1.to_excel(writer, sheet_name="Готово к импорту", index=False)
        df_sheet2.to_excel(writer, sheet_name="Сводная аналитика", index=False)
        df_sheet3.to_excel(writer, sheet_name="Детальное сопоставление", index=False)
        if not df_sheet4.empty:
            df_sheet4.to_excel(writer, sheet_name="База трекера (все товары)", index=False)

    saved_path, _ = save_excel_safely(write_workbook, output_path, "ИТОГОВЫЙ_ОТЧЕТ.xlsx")

    # Синхронизация копии во внешнюю папку если доступна
    if saved_path and os.path.exists(EXTERNAL_BACKUP_DIR):
        try:
            shutil.copy2(saved_path, os.path.join(EXTERNAL_BACKUP_DIR, "ИТОГОВЫЙ_ОТЧЕТ.xlsx"))
        except Exception:
            pass

    return saved_path


if __name__ == "__main__":
    print("=== ПРОВЕРКА МОДУЛЯ АНАЛИТИКИ ===")
    checks = check_startup_files()
    for c in checks:
        print(f"[{c['warning_level'].upper()}] {c['title']}: {c['message']}")

    locked = load_locked_articles()
    print(f"Заблокированных артикулов: {len(locked)} (Пример: {list(locked)[:5]})")

    rrc = load_rrc_prices()
    print(f"Загружено цен РРЦ: {len(rrc)}")

    cat = load_catalog_dataframe()
    print(f"Загружено позиций каталога: {len(cat)}")
